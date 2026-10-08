#!/usr/bin/env python3
"""
Generate the sample Excel workbook of query pairs from the single source of
truth (src/app/queries.py) and write it to infra/seed/queries.xlsx so Terraform
can upload it to S3.

Usage:  python3 scripts/gen_excel.py
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "src"))

from app.queries import QUERY_PAIRS  # noqa: E402

import openpyxl  # noqa: E402
from openpyxl.styles import Font, Alignment  # noqa: E402


def main() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "queries"

    headers = ["query_id", "description", "category", "postgres_sql", "mysql_sql"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")

    for q in QUERY_PAIRS:
        ws.append([q["query_id"], q["description"], q["category"],
                   q["postgres_sql"], q["mysql_sql"]])

    widths = {"A": 30, "B": 45, "C": 22, "D": 70, "E": 70}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    out_dir = os.path.join(ROOT, "infra", "seed")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "queries.xlsx")
    wb.save(out_path)
    print(f"wrote {len(QUERY_PAIRS)} query pairs -> {out_path}")


if __name__ == "__main__":
    main()
